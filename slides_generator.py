import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def create_presentation(output_pptx_path: str):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Color palette
    COLOR_BG = RGBColor(15, 23, 42)        # Slate 900
    COLOR_CARD = RGBColor(30, 41, 59)      # Slate 800
    COLOR_CARD_BORDER = RGBColor(51, 65, 85) # Slate 700
    COLOR_PRIMARY = RGBColor(56, 189, 248) # Sky 400
    COLOR_ACCENT = RGBColor(16, 185, 129)  # Emerald 500
    COLOR_AMBER = RGBColor(245, 158, 11)   # Amber 500
    COLOR_TEXT = RGBColor(248, 250, 252)   # Slate 50
    COLOR_MUTED = RGBColor(148, 163, 184)  # Slate 400

    slides_data = [
        {
            "badge": "DATA STRUCTURES & ALGORITHMS | GRAPH THEORY",
            "title": "Dijkstra's Shortest Path Algorithm",
            "subtitle": "Single-Source Shortest Path (SSSP) on Weighted Graphs",
            "author": "Robin • Grounded in GeeksforGeeks & MIT 6.006 (Demaine & Devadas)",
            "cards": [
                {
                    "title": "🎯 Core Problem",
                    "text": "Find shortest paths from a single source node to all other vertices in a directed or undirected graph with non-negative edge weights."
                },
                {
                    "title": "⚡ Greedy Paradigm",
                    "text": "Greedily finalizes the closest unvisited vertex at each step, guaranteeing optimality via mathematical edge relaxation."
                },
                {
                    "title": "🌐 Industry Applications",
                    "text": "Drives the Internet's core routing via OSPF, powers Google Maps / GPS navigation, and guides autonomous vehicle pathfinding."
                }
            ]
        },
        {
            "badge": "FOUNDATIONAL CONCEPTS",
            "title": "1. Problem Formulation & Edge Relaxation",
            "subtitle": "How Distance Estimates Converge to True Shortest Distances",
            "cards": [
                {
                    "title": "Mathematical Definition",
                    "text": "Given G = (V, E) with weights w(u, v) ≥ 0:\n• d[v]: Tentative distance estimate from source s\n• π[v]: Predecessor vertex in the Shortest Path Tree\n• δ(s, v): True minimum shortest path weight."
                },
                {
                    "title": "The Relaxation Step (Relax)",
                    "text": "if d[v] > d[u] + w(u, v):\n    d[v] = d[u] + w(u, v)\n    π[v] = u\n\nIf traversing through vertex u offers a shorter path to v, immediately update d[v]!"
                },
                {
                    "title": "Key Invariants (MIT 6.006)",
                    "text": "• Upper-Bound Property: d[v] ≥ δ(s, v) at all times\n• Triangle Inequality: δ(s, v) ≤ δ(s, u) + w(u, v)\n• Convergence: Once d[u] = δ(s, u), relaxing (u, v) sets d[v] = δ(s, v)."
                }
            ]
        },
        {
            "badge": "ALGORITHMIC MECHANICS",
            "title": "2. The Greedy Invariant & SPT Construction",
            "subtitle": "Why Greed Guarantees Global Optimality for Non-Negative Weights",
            "cards": [
                {
                    "title": "Two Vertex Partitions",
                    "text": "1. Visited / Finalized Set (S): Vertices whose shortest distance is proven optimal: d[u] = δ(s, u).\n2. Unvisited Priority Queue (Q): Vertices awaiting finalization, keyed by d[v]."
                },
                {
                    "title": "Step-by-Step Execution",
                    "text": "1. Set d[src] = 0, all other d[v] = ∞. Push (src, 0) to PQ.\n2. While PQ is not empty:\n   a. Pop vertex u with minimum d[u].\n   b. Add u to finalized set S.\n   c. For each neighbor v of u: Relax(u, v, w)."
                },
                {
                    "title": "Shortest Path Tree (SPT)",
                    "text": "The predecessor pointers π[v] form a directed spanning tree rooted at the source. Any path from s to v in this tree is a strictly shortest path."
                }
            ]
        },
        {
            "badge": "REPRESENTATION 1",
            "title": "3. Adjacency Matrix Implementation — O(V²)",
            "subtitle": "Simplicity and Optimality for Dense Graphs (GeeksforGeeks)",
            "cards": [
                {
                    "title": "How it Works",
                    "text": "• Graph stored as V x V 2D matrix.\n• To find minimum distance node: linearly scan the V-element distance array in O(V) time.\n• Outer loop executes V times → Total Time: O(V²)."
                },
                {
                    "title": "Space & Resource Trade-Off",
                    "text": "• Time Complexity: O(V²)\n• Space Complexity: O(V²) for matrix, O(V) auxiliary.\n• Best for dense graphs where E ≈ V² (e.g., airline routing between all major cities)."
                },
                {
                    "title": "Limitations",
                    "text": "For sparse networks (e.g., roads, social graphs with E = O(V)), scanning all V vertices at every step wastefully inspects non-existent edges."
                }
            ]
        },
        {
            "badge": "REPRESENTATION 2",
            "title": "4. Adjacency List + Min-Heap — O((V + E) log V)",
            "subtitle": "The Production Standard for Sparse Networks (GeeksforGeeks)",
            "cards": [
                {
                    "title": "Heap Acceleration",
                    "text": "• Store graph as adjacency list.\n• Maintain tentative distances in a Min-Heap (Priority Queue).\n• Extract-Min drops from O(V) to O(log V)!\n• Edge relaxation inserts in O(log V)."
                },
                {
                    "title": "Lazy Deletion Strategy",
                    "text": "• Instead of complex decrease-key handles, push (new_dist, v) directly into the heap.\n• When popping, if current_dist > dist[v], discard as stale!\n• Maintains identical O((V + E) log V) bound."
                },
                {
                    "title": "Quantitative Comparison",
                    "text": "For V = 100,000 and E = 300,000:\n• Matrix O(V²): 10,000,000,000 operations (~10s)\n• Min-Heap O((V+E) log V): ~6,600,000 ops (<0.01s)\n→ Over 1,500x speedup in real-world graphs!"
                }
            ]
        },
        {
            "badge": "C++ STL OPTIMIZATION",
            "title": "5. Using std::set in C++ (Self-Balancing BST)",
            "subtitle": "Emulating Decrease-Key without Heap Redundancy",
            "cards": [
                {
                    "title": "Why std::set?",
                    "text": "• std::priority_queue in C++ does not support erasing arbitrary elements.\n• std::set is implemented as a Red-Black Tree and supports erase() in O(log N)."
                },
                {
                    "title": "The Dynamic Update Mechanism",
                    "text": "When relaxing edge (u, v):\nif (dist[v] != INF)\n    st.erase(st.find({dist[v], v}));\ndist[v] = dist[u] + weight;\nst.insert({dist[v], v});\n\nOld entries are purged immediately!"
                },
                {
                    "title": "Space Guarantee",
                    "text": "The set strictly contains at most V elements at all times (O(V) space), compared to up to E elements in priority_queue lazy deletion."
                }
            ]
        },
        {
            "badge": "THEORETICAL LIMITS",
            "title": "6. Why Dijkstra Fails on Negative Weights",
            "subtitle": "The Greedy Invariant Breakdown & Counterexample",
            "cards": [
                {
                    "title": "The Fatal Flaw",
                    "text": "• Dijkstra assumes path weights monotonically increase with each hop.\n• Negative edges allow a path to become shorter after traversing more edges.\n• Once a node is marked visited, Dijkstra never re-checks it!"
                },
                {
                    "title": "The 3-Node Counterexample",
                    "text": "• Edges: (A, B)=5, (A, C)=2, (B, C)= -6\n• Dijkstra finalizes C first with dist[C] = 2.\n• Later it inspects B (5) and finds A → B → C = 5 - 6 = -1.\n• Result: Dijkstra misses the true shortest path (-1)!"
                },
                {
                    "title": "Correct Alternative",
                    "text": "• Use Bellman-Ford Algorithm: O(V · E).\n• Relaxes all E edges V-1 times and detects negative cycles.\n• For DAGs with negative weights, topological sort relaxation achieves O(V + E)!"
                }
            ]
        },
        {
            "badge": "REAL-WORLD SYSTEMS",
            "title": "7. Industrial Applications of Dijkstra",
            "subtitle": "From Internet Backbones to Planetary Navigation",
            "cards": [
                {
                    "title": "OSPF & IS-IS Internet Routing",
                    "text": "• Routers broadcast Link-State Advertisements (LSAs).\n• Every router runs Dijkstra with itself as root to build its Forwarding Information Base (FIB)."
                },
                {
                    "title": "Google Maps & GPS",
                    "text": "• Road networks mapped with weights = distance / speed + congestion.\n• Uses Bidirectional Dijkstra, A* Heuristics, and Contraction Hierarchies for sub-millisecond route calculation."
                },
                {
                    "title": "Robotics & Telecommunications",
                    "text": "• Grid maps: autonomous robots find collision-free minimum battery trajectories.\n• High-frequency trading: calculates lowest-latency fiber-optic routes between exchanges."
                }
            ]
        },
        {
            "badge": "SYSTEM ARCHITECTURE",
            "title": "8. Grounded RAG Pipeline Architecture",
            "subtitle": "Ensuring Factual Integrity & Transparent Source Attribution",
            "cards": [
                {
                    "title": "Knowledge Ingestion",
                    "text": "Authentic study archive curated from GeeksforGeeks and MIT 6.006 (Demaine & Devadas), split into semantic code & math chunks."
                },
                {
                    "title": "Hybrid Retrieval Engine",
                    "text": "Combines dense cosine similarity (semantic intent) with BM25 lexical token matching for exact algorithmic formulas and Big-O bounds."
                },
                {
                    "title": "Strict Grounding Guardrail",
                    "text": "Out-of-scope filter rejects non-Dijkstra queries, and verified answers output clickable citations with exact page numbers and URLs."
                }
            ]
        },
        {
            "badge": "CONCLUSION & SUMMARY",
            "title": "9. Key Algorithmic Insights & Takeaways",
            "subtitle": "The Definitive Guide to Dijkstra's Algorithm",
            "cards": [
                {
                    "title": "Optimal Choice",
                    "text": "• Dense graphs (E ≈ V²): Adjacency Matrix in O(V²).\n• Sparse graphs (E ≪ V²): Adjacency List + Min-Heap in O((V + E) log V)."
                },
                {
                    "title": "Critical Rule",
                    "text": "All edge weights MUST be non-negative (w(e) ≥ 0). If negative weights exist, switch to Bellman-Ford O(V · E)."
                },
                {
                    "title": "Interactive Playground",
                    "text": "Explore step-by-step edge relaxation, inspect live priority queue states, and query the grounded RAG assistant directly in the web app!"
                }
            ]
        }
    ]

    for data in slides_data:
        slide = prs.slides.add_slide(blank_layout)

        # Background fill
        bg_shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0), Inches(0), Inches(13.333), Inches(7.5))
        bg_shape.fill.solid()
        bg_shape.fill.fore_color.rgb = COLOR_BG
        bg_shape.line.color.rgb = COLOR_BG

        # Top Badge
        badge_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.5), Inches(11.7), Inches(0.4))
        tf_badge = badge_box.text_frame
        tf_badge.word_wrap = True
        p_badge = tf_badge.paragraphs[0]
        p_badge.text = data["badge"]
        p_badge.font.size = Pt(11)
        p_badge.font.bold = True
        p_badge.font.color.rgb = COLOR_PRIMARY

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.85), Inches(11.7), Inches(0.8))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = data["title"]
        p_title.font.size = Pt(28)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_TEXT

        # Subtitle
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(0.4))
        tf_sub = sub_box.text_frame
        tf_sub.word_wrap = True
        p_sub = tf_sub.paragraphs[0]
        p_sub.text = data["subtitle"]
        p_sub.font.size = Pt(14)
        p_sub.font.color.rgb = COLOR_MUTED

        # Content Cards (3 cards side by side)
        cards = data.get("cards", [])
        if cards:
            card_width = Inches(3.64)
            card_height = Inches(4.7)
            card_top = Inches(2.2)
            card_gap = Inches(0.39)

            for idx, c in enumerate(cards):
                c_left = Inches(0.8) + idx * (card_width + card_gap)
                
                # Card background shape
                card_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, c_left, card_top, card_width, card_height)
                card_shape.fill.solid()
                card_shape.fill.fore_color.rgb = COLOR_CARD
                card_shape.line.color.rgb = COLOR_CARD_BORDER
                card_shape.line.width = Pt(1)

                # Card title
                c_title_box = slide.shapes.add_textbox(c_left + Inches(0.2), card_top + Inches(0.2), card_width - Inches(0.4), Inches(0.6))
                tf_ct = c_title_box.text_frame
                tf_ct.word_wrap = True
                p_ct = tf_ct.paragraphs[0]
                p_ct.text = c["title"]
                p_ct.font.size = Pt(16)
                p_ct.font.bold = True
                p_ct.font.color.rgb = COLOR_AMBER if idx == 1 else (COLOR_ACCENT if idx == 0 else COLOR_PRIMARY)

                # Card content text
                c_body_box = slide.shapes.add_textbox(c_left + Inches(0.2), card_top + Inches(0.85), card_width - Inches(0.4), card_height - Inches(1.0))
                tf_cb = c_body_box.text_frame
                tf_cb.word_wrap = True
                p_cb = tf_cb.paragraphs[0]
                p_cb.text = c["text"]
                p_cb.font.size = Pt(12)
                p_cb.font.color.rgb = COLOR_TEXT

    prs.save(output_pptx_path)
    print(f"Presentation saved successfully to: {output_pptx_path}")

def create_html_presentation(output_html_path: str):
    html_content = """<!DOCTYPE html>
<html lang="en">
<head>
    <meta charset="UTF-8">
    <meta name="viewport" content="width=device-width, initial-scale=1.0">
    <title>Dijkstra's Shortest Path Algorithm — Master Slide Deck</title>
    <style>
        * { box-sizing: border-box; margin: 0; padding: 0; }
        body { font-family: system-ui, -apple-system, sans-serif; background: #0b0f19; color: #f8fafc; overflow: hidden; height: 100vh; display: flex; flex-direction: column; }
        header { background: #0f172a; padding: 14px 24px; border-bottom: 1px solid #1e293b; display: flex; justify-content: space-between; align-items: center; }
        .logo { font-size: 17px; font-weight: bold; color: #38bdf8; display: flex; align-items: center; gap: 8px; }
        .counter { font-size: 14px; font-family: monospace; color: #94a3b8; }
        .nav-btn { background: #1e293b; border: 1px solid #334155; color: #f8fafc; padding: 6px 14px; border-radius: 6px; cursor: pointer; font-size: 13px; font-weight: 600; }
        .nav-btn:hover { background: #334155; }
        .deck-container { flex: 1; display: flex; align-items: center; justify-content: center; padding: 24px; position: relative; }
        .slide { display: none; width: 100%; max-width: 1100px; background: #0f172a; border: 1px solid #1e293b; border-radius: 16px; padding: 40px; box-shadow: 0 20px 40px rgba(0,0,0,0.5); }
        .slide.active { display: block; animation: fadeIn 0.3s ease; }
        @keyframes fadeIn { from { opacity: 0; transform: translateY(10px); } to { opacity: 1; transform: translateY(0); } }
        .badge { display: inline-block; background: rgba(56, 189, 248, 0.1); border: 1px solid rgba(56, 189, 248, 0.3); color: #38bdf8; font-size: 11px; font-weight: 700; padding: 4px 10px; border-radius: 20px; letter-spacing: 1px; margin-bottom: 12px; }
        h1 { font-size: 28px; font-weight: 800; color: #f8fafc; margin-bottom: 8px; }
        .subtitle { font-size: 15px; color: #94a3b8; margin-bottom: 30px; }
        .cards-grid { display: grid; grid-template-columns: repeat(3, 1fr); gap: 20px; }
        .card { background: #1e293b; border: 1px solid #334155; border-radius: 12px; padding: 22px; }
        .card h3 { font-size: 16px; font-weight: 700; color: #38bdf8; margin-bottom: 12px; }
        .card p { font-size: 13px; line-height: 1.6; color: #cbd5e1; white-space: pre-line; }
        footer { background: #0f172a; padding: 12px 24px; border-top: 1px solid #1e293b; display: flex; justify-content: space-between; align-items: center; }
    </style>
</head>
<body>
    <header>
        <div class="logo">🛣️ Dijkstra's Algorithm Slide Deck</div>
        <div class="counter" id="slideNum">Slide 1 / 9</div>
        <div style="display:flex; gap:10px;">
            <button class="nav-btn" onclick="prevSlide()">◀ Previous</button>
            <button class="nav-btn" onclick="nextSlide()">Next ▶</button>
        </div>
    </header>

    <div class="deck-container">
        <!-- Slide 1 -->
        <div class="slide active">
            <div class="badge">DATA STRUCTURES & ALGORITHMS • GRAPH THEORY</div>
            <h1>Dijkstra's Shortest Path Algorithm</h1>
            <div class="subtitle">Single-Source Shortest Path (SSSP) on Weighted Graphs • Grounded in GeeksforGeeks & MIT 6.006</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>🎯 The Problem</h3>
                    <p>Find shortest paths from a single source node to all other vertices in a directed or undirected graph with non-negative edge weights.</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">⚡ Greedy Paradigm</h3>
                    <p>Greedily finalizes the closest unvisited vertex at each step, guaranteeing global optimality via edge relaxation.</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">🌐 Applications</h3>
                    <p>Drives the Internet's core routing via OSPF, powers Google Maps / GPS navigation, and guides autonomous vehicle pathfinding.</p>
                </div>
            </div>
        </div>

        <!-- Slide 2 -->
        <div class="slide">
            <div class="badge">FOUNDATIONS</div>
            <h1>1. Problem Formulation & Edge Relaxation</h1>
            <div class="subtitle">How Distance Estimates Converge to True Shortest Distances</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>Mathematical Model</h3>
                    <p>Graph G = (V, E) with w(u, v) ≥ 0:
• d[v]: Tentative shortest-path estimate
• π[v]: Predecessor vertex
• δ(s, v): True minimum distance.</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Relaxation Step</h3>
                    <p>Relax(u, v, w):
if d[v] > d[u] + w(u, v):
    d[v] = d[u] + w(u, v)
    π[v] = u

Lowers upper bound when a shorter path is found!</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">MIT 6.006 Invariants</h3>
                    <p>• Upper-Bound: d[v] ≥ δ(s, v) always
• Triangle Inequality: δ(s, v) ≤ δ(s, u) + w(u, v)
• Convergence: Once d[u] is optimal, relaxing (u, v) optimizes d[v].</p>
                </div>
            </div>
        </div>

        <!-- Slide 3 -->
        <div class="slide">
            <div class="badge">MECHANICS</div>
            <h1>2. Greedy Invariant & Shortest Path Tree</h1>
            <div class="subtitle">Why Greed Guarantees Global Optimality for Non-Negative Weights</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>Vertex Partitions</h3>
                    <p>1. Visited Set (S): Vertices whose shortest distance is finalized: d[u] = δ(s, u).
2. Priority Queue (Q): Vertices awaiting finalization, keyed by d[v].</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Algorithm Loop</h3>
                    <p>1. Initialize d[s] = 0, others = inf.
2. While Q is not empty:
   a. u = Extract-Min(Q)
   b. Add u to S.
   c. For each neighbor v: Relax(u, v, w).</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">Shortest Path Tree</h3>
                    <p>Predecessor pointers π[v] form a directed spanning tree rooted at source s. Tracing backwards yields the exact sequence of vertices.</p>
                </div>
            </div>
        </div>

        <!-- Slide 4 -->
        <div class="slide">
            <div class="badge">REPRESENTATION 1</div>
            <h1>3. Adjacency Matrix Implementation — O(V²)</h1>
            <div class="subtitle">Simplicity and Optimality for Dense Graphs (GeeksforGeeks)</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>How it Works</h3>
                    <p>• Graph stored as V x V 2D matrix.
• Finding minimum distance vertex requires a linear scan of V elements in O(V).
• Total Time: V × O(V) = O(V²).</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Space & Complexity</h3>
                    <p>• Time: O(V²)
• Space: O(V²) for matrix, O(V) auxiliary.
• Optimal for dense graphs where E ≈ V².</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">When to Avoid</h3>
                    <p>For sparse networks (roads, social graphs with E = O(V)), scanning all V vertices at every step wastefully inspects non-existent edges.</p>
                </div>
            </div>
        </div>

        <!-- Slide 5 -->
        <div class="slide">
            <div class="badge">REPRESENTATION 2</div>
            <h1>4. Adjacency List + Min-Heap — O((V + E) log V)</h1>
            <div class="subtitle">The Production Standard for Sparse Networks (GeeksforGeeks)</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>Heap Optimization</h3>
                    <p>• Graph stored as adjacency list.
• Tentative distances kept in a Min-Heap.
• Extract-Min drops from O(V) to O(log V)!
• Edge relaxation takes O(log V).</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Lazy Deletion</h3>
                    <p>• Push (new_dist, v) into the heap on relaxation.
• Discard stale entries upon popping if current_dist > dist[v].
• Maintains identical O((V + E) log V) bound!</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">Real-World Speedup</h3>
                    <p>For V = 100,000, E = 300,000:
• Matrix: 10,000,000,000 ops (~10s)
• Min-Heap: 6,600,000 ops (<0.01s)
→ Over 1,500x faster!</p>
                </div>
            </div>
        </div>

        <!-- Slide 6 -->
        <div class="slide">
            <div class="badge">C++ STL OPTIMIZATION</div>
            <h1>5. Using std::set in C++ (Balanced BST)</h1>
            <div class="subtitle">Emulating Decrease-Key without Redundant Heap Growth</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>Why std::set?</h3>
                    <p>• std::priority_queue does not support arbitrary deletion.
• std::set (Red-Black Tree) supports erase() in O(log N).</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Dynamic Update</h3>
                    <p>if (dist[v] != INF)
    st.erase(st.find({dist[v], v}));
dist[v] = dist[u] + weight;
st.insert({dist[v], v});

Purges old entries immediately!</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">Strict O(V) Space</h3>
                    <p>The set contains at most V elements at all times, preventing heap bloat up to E elements in memory-constrained devices.</p>
                </div>
            </div>
        </div>

        <!-- Slide 7 -->
        <div class="slide">
            <div class="badge">THEORETICAL LIMITS</div>
            <h1>6. Why Dijkstra Fails on Negative Weights</h1>
            <div class="subtitle">Greedy Invariant Breakdown & Counterexample</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>The Fatal Flaw</h3>
                    <p>• Assumes path weights only grow with more hops.
• Negative edges allow a path to decrease.
• Once a node is finalized, Dijkstra never re-evaluates it!</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">3-Node Counterexample</h3>
                    <p>Edges: (A, B)=5, (A, C)=2, (B, C)= -6
• Dijkstra finalizes C first with dist[C] = 2.
• Later A → B → C = 5 - 6 = -1.
• Misses the true shortest path (-1)!</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">What to Use Instead</h3>
                    <p>• Bellman-Ford: O(V · E), detects negative cycles.
• DAGs: Topological sort relaxation in O(V + E).</p>
                </div>
            </div>
        </div>

        <!-- Slide 8 -->
        <div class="slide">
            <div class="badge">APPLICATIONS</div>
            <h1>7. Industrial Applications of Dijkstra</h1>
            <div class="subtitle">Internet Backbones, GPS Routing, and Robotics</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>OSPF Network Routing</h3>
                    <p>Core Internet routers exchange Link-State Advertisements (LSAs) and run Dijkstra to build their IP Forwarding Information Base (FIB).</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Google Maps & GPS</h3>
                    <p>Intersections = nodes; roads = edges. Enhanced with A* Heuristics and Contraction Hierarchies for continent-scale instant routing.</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">Robotics & Finance</h3>
                    <p>Autonomous robots plan collision-free obstacle routes; financial carriers route packets over minimum-latency fiber optic lines.</p>
                </div>
            </div>
        </div>

        <!-- Slide 9 -->
        <div class="slide">
            <div class="badge">SUMMARY</div>
            <h1>8. Key Algorithmic Insights & Takeaways</h1>
            <div class="subtitle">The Complete Decision Guide for Engineers</div>
            <div class="cards-grid">
                <div class="card">
                    <h3>Optimal Choice</h3>
                    <p>• Dense graphs (E ≈ V²): Adjacency Matrix O(V²).
• Sparse graphs (E ≪ V²): Adjacency List + Min-Heap O((V + E) log V).</p>
                </div>
                <div class="card">
                    <h3 style="color:#f59e0b;">Cardinal Rule</h3>
                    <p>All edge weights MUST be non-negative (w(e) ≥ 0). If negative edges exist, use Bellman-Ford O(V · E).</p>
                </div>
                <div class="card">
                    <h3 style="color:#10b981;">Interactive Suite</h3>
                    <p>Explore step-by-step edge relaxation, inspect live priority queues, and query the grounded RAG assistant directly!</p>
                </div>
            </div>
        </div>
    </div>

    <footer>
        <div style="color:#64748b; font-size:12px;">Authored by Robin • Dijkstra Grounded SOT System</div>
        <div style="color:#64748b; font-size:12px;">Use Left / Right Arrow keys to navigate</div>
    </footer>

    <script>
        let currentSlide = 0;
        const slides = document.querySelectorAll('.slide');
        const numDisplay = document.getElementById('slideNum');

        function updateSlide() {
            slides.forEach((s, idx) => {
                s.classList.toggle('active', idx === currentSlide);
            });
            numDisplay.textContent = `Slide ${currentSlide + 1} / ${slides.length}`;
        }

        function nextSlide() {
            currentSlide = (currentSlide + 1) % slides.length;
            updateSlide();
        }

        function prevSlide() {
            currentSlide = (currentSlide - 1 + slides.length) % slides.length;
            updateSlide();
        }

        document.addEventListener('keydown', (e) => {
            if (e.key === 'ArrowRight' || e.key === ' ') nextSlide();
            if (e.key === 'ArrowLeft') prevSlide();
        });
    </script>
</body>
</html>
"""
    with open(output_html_path, "w", encoding="utf-8") as f:
        f.write(html_content)
    print(f"HTML slides saved successfully to: {output_html_path}")

if __name__ == "__main__":
    os.makedirs("static", exist_ok=True)
    create_presentation("static/Dijkstra_Algorithm_Complete_Guide.pptx")
    create_html_presentation("static/slides.html")
